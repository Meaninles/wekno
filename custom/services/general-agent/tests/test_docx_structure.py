from io import BytesIO
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from weknora_docreader_runtime.structure import expand_merged_docx_cells


class DocxStructureTests(unittest.TestCase):
    def round_trip(self,doc):
        data=BytesIO();doc.save(data);original=data.getvalue()
        return Document(BytesIO(expand_merged_docx_cells(original))),Document(BytesIO(original))

    def test_declared_horizontal_and_vertical_merges_only(self):
        doc=Document();table=doc.add_table(rows=3,cols=3)
        table.cell(0,0).merge(table.cell(1,0)).text="Owner"
        table.cell(0,1).merge(table.cell(0,2)).text="Header"
        table.cell(1,1).text="Value";table.cell(2,1).text="Next"
        parsed,original=self.round_trip(doc)
        self.assertEqual([c.text for c in parsed.tables[0].rows[0].cells],["Owner","Header","Header"])
        self.assertEqual([c.text for c in parsed.tables[0].rows[1].cells],["Owner","Value",""])
        self.assertEqual([c.text for c in parsed.tables[0].rows[2].cells],["","Next",""])
        self.assertTrue(original._element.xpath('.//w:vMerge | .//w:gridSpan'))
        self.assertFalse(parsed._element.xpath('.//w:vMerge | .//w:gridSpan'))

    def test_nested_table_and_omitted_grid_positions(self):
        doc=Document();outer=doc.add_table(rows=1,cols=2)
        nested=outer.cell(0,0).add_table(rows=2,cols=2)
        nested.cell(0,0).merge(nested.cell(1,0)).text="Nested owner"
        outer.cell(0,0).merge(outer.cell(0,1))
        grid=doc.add_table(rows=2,cols=3)
        grid.cell(0,1).merge(grid.cell(1,1)).text="Offset owner"
        tr=grid.rows[1]._tr;tr.remove(tr.tc_lst[0])
        before=OxmlElement('w:gridBefore');before.set(qn('w:val'),'1');tr.get_or_add_trPr().append(before)
        parsed,_=self.round_trip(doc)
        self.assertFalse(parsed._element.xpath('.//w:vMerge | .//w:gridSpan'))
        self.assertEqual(parsed.tables[0].cell(0,1).tables[0].cell(1,0).text,"Nested owner")
        row=parsed.tables[1].rows[1]
        self.assertEqual(row.grid_cols_before,1)
        self.assertEqual([c.text for c in row.cells],["Offset owner",""])

if __name__=='__main__':unittest.main()
