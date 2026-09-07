"""Verify the production workspace adapter inside the existing dev runtime."""
import asyncio
import json
import time
import uuid

import aiodocker
from app.contracts import RunRequest
from app.tools import invocation_id
from app.workspace import SandboxBackend


class Control:
    async def post(self, path):
        assert path == "runs/heartbeat"


async def main():
    payload = RunRequest(run_id="development-probe-"+uuid.uuid4().hex, owner_epoch=1,
        session_id="component-verification", assistant_message_id="component-verification",
        query="", llm={"model_name":"unused"}, deadline_unix=time.time()+120,
        tool_callback_url="http://unused/tools/call")
    backend = SandboxBackend(payload,Control())
    try:
        content = "中文原件，数值 42。".encode()
        await backend.write_file("inputs/原件.txt",content)
        assert await backend.read_file("inputs/原件.txt") == content
        result = await backend.exec_shell(["python","-c","import os; print(os.getuid()); assert not any('API_KEY' in k for k in os.environ)"])
        assert result.exit_code == 0 and result.stdout.strip() == b"1000"
        assert (await backend.exec_shell("printf '%s' 'shell works'")).stdout == b"shell works"
        marker = invocation_id.set("stable-execution")
        try:
            command = ["python","-c","with open('once.txt','a') as f: f.write('once')"]
            for _ in range(2):
                assert (await backend.exec_shell(command)).exit_code == 0
            assert await backend.read_file("once.txt") == b"once"
        finally:
            invocation_id.reset(marker)
        result = await backend.exec_shell(["python","-c","print('x' * (10*1024*1024))"])
        assert result.exit_code == 0 and len(result.stdout) < 1024**2+150
        assert (await backend.exec_shell(["ln","-s","/control","/workspace/escape"])).exit_code == 0
        for operation in (backend.write_file("escape/forbidden",b"bad"), backend.read_file("escape/epoch")):
            try:
                await operation
                raise AssertionError("Workspace followed a private-directory symlink")
            except OSError:
                pass
        assert (await backend.exec_shell(["mkfifo", "/workspace/pipe"])).exit_code == 0
        try:
            await asyncio.wait_for(backend.read_file("pipe"), timeout=3)
            raise AssertionError("Workspace read a FIFO as a regular file")
        except OSError:
            pass
        office = """from docx import Document
from openpyxl import Workbook
from pptx import Presentation
doc=Document(); doc.add_paragraph('统一内核文件验证'); doc.save('check.docx')
book=Workbook(); book.active['A1']=42; book.save('check.xlsx')
slides=Presentation(); slides.slides.add_slide(slides.slide_layouts[0]); slides.save('check.pptx')
"""
        result = await backend.exec_shell(["python","-c",office])
        assert result.exit_code == 0, result.stderr.decode()
        for name in ("check.docx","check.xlsx","check.pptx"):
            assert (await backend.read_file(name)).startswith(b"PK")
        result = await backend.exec_shell(["libreoffice","-env:UserInstallation=file:///tmp/lo-probe","--headless","--convert-to","pdf","--outdir","/workspace","/workspace/check.docx"],timeout=60)
        assert result.exit_code == 0, result.stderr.decode()
        assert (await backend.read_file("check.pdf")).startswith(b"%PDF")
        result = await backend.exec_shell(["python","-c","import time; time.sleep(30)"],timeout=.2)
        assert result.exit_code == 124
        task = asyncio.create_task(backend.exec_shell(["python","-c","import time; time.sleep(30)"]))
        await asyncio.sleep(.2)
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass
        assert backend.container is None
        print(json.dumps({"workspace":"passed","checks":["unicode bytes","unprivileged execution","no credentials",
            "execution receipt replay","bounded logs","symlink confinement","FIFO rejection","docx/xlsx/pptx","PDF conversion","timeout","cancellation"]}))
    finally:
        await backend.close()
        async with aiodocker.Docker() as docker:
            for prefix in ("agent-workspace-","agent-receipts-"):
                volume = await docker.volumes.get(prefix+backend.key)
                await volume.delete()


asyncio.run(main())
