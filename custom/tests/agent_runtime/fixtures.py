"""Synthetic files for current-development upload and interpretation probes."""
from pathlib import Path
import csv
import io
import json
import random

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from pptx import Presentation
from pptx.util import Inches
from reportlab import rl_config
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT=Path(".local-data/agent-runtime-validation/fixtures")
ROOT.mkdir(parents=True,exist_ok=True)
TEXT="Project LANTERN-7429. Owner: Chen. Budget: 186400 CNY. Deadline: 2026-11-19. Region: Guizhou."
FONT=ImageFont.truetype("C:/Windows/Fonts/arial.ttf",42)


def small():
    for suffix in ("txt","md"):
        (ROOT/("project."+suffix)).write_text(TEXT+"\nItem A: quantity 3, unit price 12. Item B: quantity 2, unit price 8. Total: 52.\n",encoding="utf-8")
    (ROOT/"project.json").write_text(json.dumps({"project":"LANTERN-7429","budget":186400,"deadline":"2026-11-19"}),encoding="utf-8")
    (ROOT/"project.html").write_text("<html><body><h1>Project facts</h1><p>"+TEXT+"</p></body></html>",encoding="utf-8")
    with (ROOT/"sales.csv").open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f);w.writerows([["item","quantity","unit_price"],["A",3,12],["B",2,8]])
    book=Workbook();sheet=book.active;sheet.title="Items"
    for row in [["Item","Quantity","Unit price","Amount"],["A",3,12,"=B2*C2"],["B",2,8,"=B3*C3"],["Total",None,None,"=SUM(D2:D3)"]]:sheet.append(row)
    book.save(ROOT/"sales.xlsx")
    doc=Document();doc.add_heading("Project brief",0);doc.add_paragraph(TEXT)
    table=doc.add_table(rows=1,cols=3)
    for cell,value in zip(table.rows[0].cells,["Item","Quantity","Unit price"]):cell.text=value
    for row in [["A","3","12"],["B","2","8"]]:
        for cell,value in zip(table.add_row().cells,row):cell.text=value
    doc.save(ROOT/"project.docx")
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text="Project brief";slide.placeholders[1].text=TEXT
    deck.save(ROOT/"project.pptx")
    pdf=canvas.Canvas(str(ROOT/"project.pdf"));pdf.drawString(36,790,TEXT);pdf.drawString(36,760,"Item A: 3 x 12 = 36; Item B: 2 x 8 = 16; Total = 52.");pdf.save()
    img=Image.new("RGB",(1600,650),"white");draw=ImageDraw.Draw(img)
    for i,line in enumerate(["Project LANTERN-7429","Budget: 186400 CNY","Deadline: 2026-11-19","A: 3 x 12 = 36","B: 2 x 8 = 16","Total: 52"]):draw.text((40,25+i*95),line,font=FONT,fill="black")
    img.save(ROOT/"facts.png");img.save(ROOT/"facts.jpg",quality=95)


def large():
    # Real embedded RGB figures dominate size, while searchable text and cells
    # remain inspectable. Both files are ~99% of the actual 128 MiB limit.
    rl_config.useA85=0
    pdf=canvas.Canvas(str(ROOT/"large-near-limit.pdf"),pageCompression=1)
    book=Workbook();book.remove(book.active)
    images=[]
    for i in range(10):
        pixels=random.Random(7429+i).randbytes(2048*2290*3)
        image=Image.frombytes("RGB",(2048,2290),pixels)
        draw=ImageDraw.Draw(image);draw.rectangle((50,50,1800,190),fill="white")
        draw.text((80,80),f"LANTERN FIGURE {i+1:02d} / VALUE {1000+i*7}",font=FONT,fill="black")
        encoded=io.BytesIO();image.save(encoded,format="PNG");encoded.seek(0);images.append(encoded)
        pdf.drawString(25,815,f"LANTERN large file page {i+1}. Control value: {1000+i*7}.")
        pdf.drawImage(ImageReader(image),25,80,width=540,height=700);pdf.showPage()
        sheet=book.create_sheet(f"Region-{i+1}");sheet.append(["Project","Region","Control value","Budget"])
        sheet.append(["LANTERN-7429",i+1,1000+i*7,186400+i*100])
        picture=SheetImage(encoded);picture.width=540;picture.height=700;sheet.add_image(picture,"A6")
    pdf.save();book.save(ROOT/"large-near-limit.xlsx")
    for name in ("large-near-limit.pdf","large-near-limit.xlsx"):
        size=(ROOT/name).stat().st_size
        assert 120*1024**2<size<128*1024**2,(name,size)
        print(json.dumps({"file":name,"bytes":size,"limit_percent":round(size/(128*1024**2)*100,2)}),flush=True)


if __name__=="__main__":
    import sys
    small()
    if "--large" in sys.argv:large()
    print(str(ROOT.resolve()))

