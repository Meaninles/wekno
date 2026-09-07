import io
import unittest
from unittest.mock import patch
from docx import Document
from docreader.parser.doc_parser import DocParser

class TestConvertedDocFormat(unittest.TestCase):
    def test_converted_doc_preserves_adjacent_numeric_cells(self):
        doc = Document()
        doc.add_paragraph("Order table")
        table = doc.add_table(rows=3, cols=3)
        for row, values in zip(table.rows, [("Item", "Quantity", "Unit price"), ("A", "3", "12"), ("B", "2", "8")]):
            for cell, value in zip(row.cells, values): cell.text = value
        content = io.BytesIO(); doc.save(content)
        parser = DocParser(file_name="order.doc", file_type="doc")
        with patch.object(parser, "_try_convert_doc_to_docx", return_value=content.getvalue()):
            text = parser._parse_with_docx("unused.doc").content
        self.assertIn("| A | 3 | 12 |", text)
        self.assertIn("| B | 2 | 8 |", text)
        self.assertNotIn("A312B28", text)
